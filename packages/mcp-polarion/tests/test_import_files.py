"""Tests for Word and ReqIF import parsing."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches

from polarion_client import import_files
from polarion_client.errors import PolarionError
from polarion_client.import_files import (
    MAX_IMPORT_BYTES,
    ImportPreview,
    parse_import_file,
)

# 1x1 PNG
_TINY_PNG = bytes(
    [
        0x89,
        0x50,
        0x4E,
        0x47,
        0x0D,
        0x0A,
        0x1A,
        0x0A,
        0x00,
        0x00,
        0x00,
        0x0D,
        0x49,
        0x48,
        0x44,
        0x52,
        0x00,
        0x00,
        0x00,
        0x01,
        0x00,
        0x00,
        0x00,
        0x01,
        0x08,
        0x06,
        0x00,
        0x00,
        0x00,
        0x1F,
        0x15,
        0xC4,
        0x89,
        0x00,
        0x00,
        0x00,
        0x0A,
        0x49,
        0x44,
        0x41,
        0x54,
        0x78,
        0x9C,
        0x63,
        0x00,
        0x01,
        0x00,
        0x00,
        0x05,
        0x00,
        0x01,
        0x0D,
        0x0A,
        0x2D,
        0xB4,
        0x00,
        0x00,
        0x00,
        0x00,
        0x49,
        0x45,
        0x4E,
        0x44,
        0xAE,
        0x42,
        0x60,
        0x82,
    ]
)

MINIMAL_REQIF = """<?xml version="1.0" encoding="UTF-8"?>
<REQ-IF xmlns="http://www.omg.org/spec/ReqIF/20110401">
  <CORE-CONTENT>
    <REQ-IF-CONTENT>
      <DATATYPES>
        <DATATYPE-DEFINITION-STRING IDENTIFIER="dt-string"/>
      </DATATYPES>
      <SPEC-TYPES>
        <SPEC-OBJECT-TYPE IDENTIFIER="sot-1" LONG-NAME="Requirement"/>
        <SPEC-RELATION-TYPE IDENTIFIER="rel-custom" LONG-NAME="CustomTrace"/>
        <ATTRIBUTE-DEFINITION-STRING IDENTIFIER="ad-name">
          <LONG-NAME>ReqIF.Name</LONG-NAME>
        </ATTRIBUTE-DEFINITION-STRING>
        <ATTRIBUTE-DEFINITION-STRING IDENTIFIER="ad-text">
          <LONG-NAME>ReqIF.Text</LONG-NAME>
        </ATTRIBUTE-DEFINITION-STRING>
        <ATTRIBUTE-DEFINITION-STRING IDENTIFIER="ad-cust">
          <LONG-NAME>CustomerID</LONG-NAME>
        </ATTRIBUTE-DEFINITION-STRING>
      </SPEC-TYPES>
      <SPEC-OBJECTS>
        <SPEC-OBJECT IDENTIFIER="obj-parent">
          <TYPE><SPEC-OBJECT-TYPE-REF>sot-1</SPEC-OBJECT-TYPE-REF></TYPE>
          <VALUES>
            <ATTRIBUTE-VALUE-STRING>
              <DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>ad-name</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION>
              <THE-VALUE>Parent requirement</THE-VALUE>
            </ATTRIBUTE-VALUE-STRING>
            <ATTRIBUTE-VALUE-STRING>
              <DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>ad-text</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION>
              <THE-VALUE><div xmlns="http://www.w3.org/1999/xhtml">Parent <b>body</b></div></THE-VALUE>
            </ATTRIBUTE-VALUE-STRING>
          </VALUES>
        </SPEC-OBJECT>
        <SPEC-OBJECT IDENTIFIER="obj-child">
          <TYPE><SPEC-OBJECT-TYPE-REF>sot-1</SPEC-OBJECT-TYPE-REF></TYPE>
          <VALUES>
            <ATTRIBUTE-VALUE-STRING>
              <DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>ad-name</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION>
              <THE-VALUE>Child requirement</THE-VALUE>
            </ATTRIBUTE-VALUE-STRING>
            <ATTRIBUTE-VALUE-STRING>
              <DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>ad-cust</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION>
              <THE-VALUE>CUST-42</THE-VALUE>
            </ATTRIBUTE-VALUE-STRING>
          </VALUES>
        </SPEC-OBJECT>
      </SPEC-OBJECTS>
      <SPEC-RELATIONS>
        <SPEC-RELATION IDENTIFIER="sr-1">
          <TYPE><SPEC-RELATION-TYPE-REF>rel-custom</SPEC-RELATION-TYPE-REF></TYPE>
          <SOURCE><SPEC-OBJECT-REF>obj-parent</SPEC-OBJECT-REF></SOURCE>
          <TARGET><SPEC-OBJECT-REF>obj-child</SPEC-OBJECT-REF></TARGET>
        </SPEC-RELATION>
      </SPEC-RELATIONS>
      <SPECIFICATIONS>
        <SPECIFICATION IDENTIFIER="spec-1" LONG-NAME="Test Specification">
          <CHILDREN>
            <SPEC-HIERARCHY IDENTIFIER="h-parent">
              <OBJECT><SPEC-OBJECT-REF>obj-parent</SPEC-OBJECT-REF></OBJECT>
              <CHILDREN>
                <SPEC-HIERARCHY IDENTIFIER="h-child">
                  <OBJECT><SPEC-OBJECT-REF>obj-child</SPEC-OBJECT-REF></OBJECT>
                </SPEC-HIERARCHY>
              </CHILDREN>
            </SPEC-HIERARCHY>
          </CHILDREN>
        </SPECIFICATION>
      </SPECIFICATIONS>
    </REQ-IF-CONTENT>
  </CORE-CONTENT>
</REQ-IF>
"""


def _build_sample_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Top requirement", level=1)
    p = doc.add_paragraph()
    run = p.add_run("Bold detail")
    run.bold = True
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell A"
    doc.add_picture(io.BytesIO(_TINY_PNG), width=Inches(0.1))
    doc.save(str(path))


def test_parse_docx_work_item_and_content(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)

    preview = parse_import_file(docx_path)
    assert len(preview.documents) == 1
    doc = preview.documents[0]
    assert len(doc.work_items) == 1
    wi = doc.work_items[0]
    assert wi.title == "Top requirement"
    assert "<strong>" in wi.description_html or "<b>" in wi.description_html.lower()
    assert "Bold detail" in wi.description_html
    assert "<table>" in wi.description_html
    assert "Cell A" in wi.description_html
    assert wi.level == 1
    assert wi.parent_key is None
    assert len(doc.images) >= 1 or any("image" in w.lower() for w in doc.warnings)


def test_parse_reqif_hierarchy_and_attributes(tmp_path: Path) -> None:
    reqif_path = tmp_path / "sample.reqif"
    reqif_path.write_text(MINIMAL_REQIF, encoding="utf-8")

    preview = parse_import_file(
        reqif_path,
        structure_link_role="has_parent",
    )
    doc = preview.documents[0]
    assert doc.module_name
    assert len(doc.work_items) == 2
    parent = doc.work_items[0]
    child = doc.work_items[1]
    assert parent.key == "obj-parent"
    assert child.key == "obj-child"
    assert child.parent_key == "obj-parent"
    assert child.reqif_id == "obj-child"
    assert child.extra_attributes.get("CustomerID") == "CUST-42"
    assert "CustomerID" in child.description_html
    hierarchy_links = [
        lnk
        for lnk in doc.links
        if lnk.relation_type == "parent" and lnk.role == "has_parent"
    ]
    assert any(lnk.source_key == "obj-parent" and lnk.target_key == "obj-child" for lnk in hierarchy_links)
    assert any("Unmapped relation" in w for w in doc.warnings)


def test_reject_missing_unsupported_oversized(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(PolarionError, match="not found"):
        parse_import_file(tmp_path / "missing.docx")

    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"fake")
    with pytest.raises(PolarionError, match="Unsupported"):
        parse_import_file(legacy)

    big = tmp_path / "big.docx"
    _build_sample_docx(big)
    monkeypatch.setattr(import_files, "MAX_IMPORT_BYTES", 10)
    with pytest.raises(PolarionError, match="maximum size"):
        parse_import_file(big)


def test_import_preview_to_dict_excludes_image_bytes(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    preview: ImportPreview = parse_import_file(docx_path)
    data = preview.to_dict()
    for doc in data["documents"]:
        for img in doc["images"]:
            assert "data" not in img
            assert "size" in img
            assert "filename" in img
